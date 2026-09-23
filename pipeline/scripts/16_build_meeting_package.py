#!/usr/bin/env python3
"""Build a radiologist/professor meeting package from the 83 normal CT scans.

The package uses the study-specific age buckets from normal_cases_age.csv,
retains every detected AirMorph generation, creates paper-style central-airway
tables, divides the full segmented tracheal centerline into distal/middle/
proximal thirds, and supplies lung-volume-normalized sensitivity analyses.
"""
from __future__ import annotations

import argparse
import math
import shutil
import tarfile
from collections import defaultdict
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import SimpleITK as sitk
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy import ndimage, sparse
from scipy.sparse.csgraph import dijkstra, minimum_spanning_tree


RUN = Path(__file__).resolve().parents[1]
DEFAULT_OUT = RUN / "AIRWAY_RADIOLOGIST_PROF_MEETING_2026-09-11"
BRANCHES = RUN / "normal_airway_branch_morphometry.csv"
AGES = RUN / "normal_cases_age.csv"
STATUS = RUN / "normal_airway_morphometry_status.csv"
MAX_REPORTED_GENERATION = None

AGE_ORDER = ["0-3", "3-6", "6-9", "10-12", "12-14", "14-16", "16-18"]
AGE_DISPLAY = {
    "0-3": "0–3",
    "3-6": "3–6",
    "6-9": "6–9",
    "10-12": "10–12",
    "12-14": "12–14",
    "14-16": "14–16",
    "16-18": "16–18",
    "All": "All",
}

NEIGHBOR_OFFSETS = [
    (dz, dy, dx)
    for dz in (-1, 0, 1)
    for dy in (-1, 0, 1)
    for dx in (-1, 0, 1)
    if (dz, dy, dx) > (0, 0, 0)
]

CENTRAL_METRICS = [
    "full_segmented_tracheal_length_mm",
    "tracheal_edt_mean_diameter_mm",
    "lmb_edt_mean_diameter_mm",
    "rmb_edt_mean_diameter_mm",
]
SECTION_METRICS = [
    "trachea_distal_edt_mean_diameter_mm",
    "trachea_middle_edt_mean_diameter_mm",
    "trachea_proximal_edt_mean_diameter_mm",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--reuse-case-cache", action="store_true")
    return parser.parse_args()


def descendants(start: int, children: dict[int, list[int]]) -> list[int]:
    result, stack = [], [start]
    while stack:
        branch_id = stack.pop()
        result.append(branch_id)
        stack.extend(children.get(branch_id, []))
    return result


def side_votes(row) -> tuple[int, int]:
    left = right = 0
    for column in ("branch_name", "lobar_name", "segmental_name", "subsegmental_name"):
        value = str(getattr(row, column, ""))
        left += int(value.startswith(("Left", "LB")))
        right += int(value.startswith(("Right", "RB")))
    return left, right


def branch_physical_x(row, case_id: str) -> float:
    """Return a branch centroid's patient-left coordinate for label fallback."""
    case_dir = Path(row.source_mask).parent
    parsing_image = sitk.ReadImage(str(case_dir / f"{case_id}_skel_parsing.nii.gz"))
    parsing = sitk.GetArrayFromImage(parsing_image)
    branch_id = int(row.branch_instance_id)
    coordinates = np.argwhere(parsing == branch_id)
    if not len(coordinates):
        raise ValueError(f"{case_id}: branch {branch_id} absent from skeleton parsing")
    z, y, x = coordinates.mean(axis=0)
    return float(
        parsing_image.TransformContinuousIndexToPhysicalPoint(
            (float(x), float(y), float(z))
        )[0]
    )


def extract_central_topology(branches: pd.DataFrame) -> pd.DataFrame:
    records = []
    for case_id, group in branches.groupby("case_id", sort=True):
        rows = {int(row.branch_instance_id): row for row in group.itertuples()}
        children: dict[int, list[int]] = defaultdict(list)
        roots = []
        for row in group.itertuples():
            branch_id = int(row.branch_instance_id)
            if pd.isna(row.parent_instance_id):
                roots.append(branch_id)
            else:
                children[int(row.parent_instance_id)].append(branch_id)
        if len(roots) != 1:
            raise ValueError(f"{case_id}: expected one tracheal root, found {len(roots)}")
        root_id = roots[0]

        candidates = []
        for child_id in children.get(root_id, []):
            subtree = descendants(child_id, children)
            left = right = 0
            for branch_id in subtree:
                l_votes, r_votes = side_votes(rows[branch_id])
                left += l_votes
                right += r_votes
            candidates.append(
                {
                    "branch_id": child_id,
                    "descendant_count": len(subtree),
                    "left_votes": left,
                    "right_votes": right,
                }
            )
        candidates.sort(key=lambda item: item["descendant_count"], reverse=True)
        selected = candidates[:2]
        if len(selected) != 2:
            raise ValueError(f"{case_id}: fewer than two root children")
        for item in selected:
            item["side"] = "L" if item["left_votes"] > item["right_votes"] else "R"
        assignment_method = "anatomic-label-votes"
        if {item["side"] for item in selected} != {"L", "R"}:
            # Sparse pediatric trees can give both main bronchi the same predicted
            # name. In LPS physical coordinates, increasing x points to patient left.
            for item in selected:
                item["physical_x"] = branch_physical_x(
                    rows[item["branch_id"]], case_id
                )
            selected.sort(key=lambda item: item["physical_x"])
            selected[0]["side"] = "R"
            selected[1]["side"] = "L"
            assignment_method = "physical-x-fallback"
        by_side = {item["side"]: item for item in selected}

        root = rows[root_id]
        lmb = rows[by_side["L"]["branch_id"]]
        rmb = rows[by_side["R"]["branch_id"]]
        records.append(
            {
                "case_id": case_id,
                "subject": root.subject,
                "root_branch_id": root_id,
                "lmb_branch_id": int(lmb.branch_instance_id),
                "rmb_branch_id": int(rmb.branch_instance_id),
                "raw_root_mst_length_mm": float(root.centerline_length_mm),
                "raw_root_chord_length_mm": float(root.chord_length_mm),
                "raw_root_edt_mean_diameter_mm": float(root.diameter_mean_mm),
                "lmb_edt_mean_diameter_mm": float(lmb.diameter_mean_mm),
                "rmb_edt_mean_diameter_mm": float(rmb.diameter_mean_mm),
                "lmb_descendant_count": by_side["L"]["descendant_count"],
                "rmb_descendant_count": by_side["R"]["descendant_count"],
                "lmb_side_vote_margin": by_side["L"]["left_votes"] - by_side["L"]["right_votes"],
                "rmb_side_vote_margin": by_side["R"]["right_votes"] - by_side["R"]["left_votes"],
                "main_bronchus_assignment_method": assignment_method,
            }
        )
    return pd.DataFrame(records)


def build_tree_path(coords: np.ndarray, spacing_zyx: np.ndarray):
    index = {tuple(coord): i for i, coord in enumerate(coords)}
    rows, cols, weights = [], [], []
    for i, coord in enumerate(coords):
        for offset in NEIGHBOR_OFFSETS:
            j = index.get(tuple(coord + offset))
            if j is None:
                continue
            weight = float(np.linalg.norm(np.asarray(offset) * spacing_zyx))
            rows.extend((i, j))
            cols.extend((j, i))
            weights.extend((weight, weight))
    graph = sparse.csr_matrix((weights, (rows, cols)), shape=(len(coords), len(coords)))
    tree = minimum_spanning_tree(graph)
    undirected = (tree + tree.T).tocsr()
    degree = np.asarray(undirected.astype(bool).sum(axis=1)).ravel()
    endpoints = np.flatnonzero(degree <= 1)
    if len(endpoints) < 2:
        raise ValueError("root skeleton has fewer than two tree endpoints")
    distances, predecessors = dijkstra(
        undirected, directed=False, indices=endpoints, return_predecessors=True
    )
    endpoint_distances = distances[:, endpoints]
    source_row, target_col = np.unravel_index(
        np.nanargmax(endpoint_distances), endpoint_distances.shape
    )
    source = int(endpoints[source_row])
    target = int(endpoints[target_col])
    path = [target]
    cursor = target
    while cursor != source:
        cursor = int(predecessors[source_row, cursor])
        if cursor < 0:
            raise ValueError("could not reconstruct root tree path")
        path.append(cursor)
    path.reverse()
    path = np.asarray(path, dtype=int)
    path_coords = coords[path]
    steps = np.linalg.norm(np.diff(path_coords, axis=0) * spacing_zyx, axis=1)
    cumulative = np.r_[0.0, np.cumsum(steps)]
    return path_coords, cumulative, len(endpoints), float(tree.sum())


def physical_z(image: sitk.Image, coord_zyx: np.ndarray) -> float:
    index_xyz = tuple(int(v) for v in coord_zyx[::-1])
    return float(image.TransformIndexToPhysicalPoint(index_xyz)[2])


def sampled_section_metrics(
    mask: np.ndarray,
    root_coords: np.ndarray,
    spacing_zyx: np.ndarray,
    path_coords: np.ndarray,
    cumulative: np.ndarray,
) -> dict:
    margin_voxels = np.ceil(30.0 / spacing_zyx).astype(int)
    lower = np.maximum(root_coords.min(axis=0) - margin_voxels, 0)
    upper = np.minimum(root_coords.max(axis=0) + margin_voxels + 1, mask.shape)
    slices = tuple(slice(int(lo), int(hi)) for lo, hi in zip(lower, upper))
    distance = ndimage.distance_transform_edt(mask[slices], sampling=spacing_zyx)
    relative_path = path_coords - lower
    diameter_at_path = 2.0 * distance[tuple(relative_path.T)]

    length = float(cumulative[-1])
    positions = np.arange(0.0, length, 1.0)
    if not len(positions) or positions[-1] < length:
        positions = np.r_[positions, length]
    sampled = np.interp(positions, cumulative, diameter_at_path)
    thirds = {
        "distal": (0.0, length / 3.0),
        "middle": (length / 3.0, 2.0 * length / 3.0),
        "proximal": (2.0 * length / 3.0, length + 1e-9),
    }
    result = {
        "tracheal_regular_samples": len(sampled),
        "tracheal_regular_sample_step_mm": 1.0,
        "tracheal_edt_mean_diameter_mm": float(np.mean(sampled)),
        "tracheal_edt_sd_diameter_mm": float(np.std(sampled, ddof=1)) if len(sampled) > 1 else 0.0,
        "tracheal_edt_median_diameter_mm": float(np.median(sampled)),
        "tracheal_edt_q1_diameter_mm": float(np.quantile(sampled, 0.25)),
        "tracheal_edt_q3_diameter_mm": float(np.quantile(sampled, 0.75)),
    }
    for name, (low, high) in thirds.items():
        select = (positions >= low) & (positions < high)
        values = sampled[select]
        result[f"trachea_{name}_edt_mean_diameter_mm"] = float(np.mean(values))
        result[f"trachea_{name}_edt_sd_diameter_mm"] = (
            float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
        )
        result[f"trachea_{name}_sample_count"] = int(len(values))
    return result


def measure_imaging_case(row) -> dict:
    case_id = row.case_id
    case_dir = Path(row.case_dir)
    root_id = int(row.root_branch_id)

    parse_image = sitk.ReadImage(str(case_dir / f"{case_id}_skel_parsing.nii.gz"))
    parse_array = sitk.GetArrayFromImage(parse_image)
    root_coords = np.argwhere(parse_array == root_id)
    if not len(root_coords):
        raise ValueError(f"{case_id}: root label {root_id} absent from skeleton parsing")
    spacing_zyx = np.asarray(parse_image.GetSpacing(), dtype=float)[::-1]
    path_coords, cumulative, endpoint_count, mst_total = build_tree_path(
        root_coords, spacing_zyx
    )
    # LPS physical z increases superiorly.  Orient the path carina -> superior end.
    if physical_z(parse_image, path_coords[0]) > physical_z(parse_image, path_coords[-1]):
        path_coords = path_coords[::-1]
        cumulative = cumulative[-1] - cumulative[::-1]

    mask_image = sitk.ReadImage(str(case_dir / "union3_lcc.nii.gz"))
    mask = sitk.GetArrayFromImage(mask_image) > 0
    if mask.shape != parse_array.shape:
        raise ValueError(f"{case_id}: mask and parsing geometry differ")
    section = sampled_section_metrics(
        mask, root_coords, spacing_zyx, path_coords, cumulative
    )

    lung_image = sitk.ReadImage(str(case_dir / "lunglobe.nii.gz"))
    lung = sitk.GetArrayFromImage(lung_image)
    voxel_mm3 = float(np.prod(lung_image.GetSpacing()))
    total_lung_ml = float(np.count_nonzero(lung) * voxel_mm3 / 1000.0)
    if total_lung_ml <= 0:
        raise ValueError(f"{case_id}: zero segmented lung volume")
    result = {
        "case_id": case_id,
        "full_segmented_tracheal_length_mm": float(cumulative[-1]),
        "root_tree_total_length_mm": mst_total,
        "root_tree_endpoint_count": endpoint_count,
        "root_centerline_voxels": len(root_coords),
        "total_segmented_lung_volume_mL": total_lung_ml,
        "lung_label_1_volume_mL": float(np.count_nonzero(lung == 1) * voxel_mm3 / 1000.0),
        "lung_label_2_volume_mL": float(np.count_nonzero(lung == 2) * voxel_mm3 / 1000.0),
    }
    result.update(section)
    return result


def build_case_table(
    branches: pd.DataFrame, age: pd.DataFrame, status: pd.DataFrame, cache: Path, reuse: bool
) -> pd.DataFrame:
    topology = extract_central_topology(branches)
    age = age.copy()
    if "case_id" not in age.columns:
        age["case_id"] = age["anon_subject"] + "_" + age["anon_accession"]
    needed_columns = ["case_id", "age_years", "age_bucket"]
    if "sex" in age.columns:
        needed_columns.append("sex")
    needed = age[needed_columns]
    if needed.age_bucket.isna().any():
        raise ValueError("one or more cases lack the requested age bucket")
    table = topology.merge(needed, on="case_id", validate="one_to_one")
    table = table.merge(status[["case_id", "case_dir"]], on="case_id", validate="one_to_one")

    if reuse and cache.exists():
        imaging = pd.read_csv(cache)
    else:
        imaging_rows = []
        for number, row in enumerate(table.itertuples(), 1):
            print(f"[{number:02d}/{len(table)}] imaging measurements: {row.case_id}", flush=True)
            imaging_rows.append(measure_imaging_case(row))
        imaging = pd.DataFrame(imaging_rows)
        cache.parent.mkdir(parents=True, exist_ok=True)
        imaging.to_csv(cache, index=False)
    table = table.merge(imaging, on="case_id", validate="one_to_one")

    lung_liters = table.total_segmented_lung_volume_mL / 1000.0
    table["lung_linear_scale_L_cuberoot"] = np.cbrt(lung_liters)
    for metric in CENTRAL_METRICS + SECTION_METRICS:
        table[f"{metric}_per_L_cuberoot"] = (
            table[metric] / table.lung_linear_scale_L_cuberoot
        )
    table["trachea_path_vs_raw_mst_difference_mm"] = (
        table.full_segmented_tracheal_length_mm - table.raw_root_mst_length_mm
    )
    return table


def describe(values: pd.Series) -> dict:
    values = pd.to_numeric(values, errors="coerce").dropna()
    if not len(values):
        return {key: np.nan for key in ("n", "mean", "sd", "median", "q1", "q3", "iqr", "min", "max")}
    q1, median, q3 = values.quantile([0.25, 0.5, 0.75])
    return {
        "n": int(len(values)),
        "mean": float(values.mean()),
        "sd": float(values.std(ddof=1)) if len(values) > 1 else 0.0,
        "median": float(median),
        "q1": float(q1),
        "q3": float(q3),
        "iqr": float(q3 - q1),
        "min": float(values.min()),
        "max": float(values.max()),
    }


def grouped_descriptive(
    data: pd.DataFrame, metrics: list[str], include_all: bool = True
) -> pd.DataFrame:
    records = []
    groups = [(bucket, data[data.age_bucket == bucket]) for bucket in AGE_ORDER]
    if include_all:
        groups.append(("All", data))
    for bucket, group in groups:
        for metric in metrics:
            record = {
                "age_group": AGE_DISPLAY[bucket],
                "age_bucket_code": bucket,
                "metric": metric,
                "n_unique_subjects": int(group.subject.nunique()),
                "n_scans": int(group.case_id.nunique()),
            }
            record.update(describe(group[metric]))
            records.append(record)
    return pd.DataFrame(records)


def fmt_age(values: pd.Series) -> str:
    stat = describe(values)
    return f"{stat['median']:.1f} ({stat['iqr']:.1f})"


def fmt_median_quartiles(values: pd.Series) -> str:
    stat = describe(values)
    return f"{stat['median']:.1f} ({stat['q1']:.1f}, {stat['q3']:.1f})"


def fmt_median_iqr(values: pd.Series) -> str:
    stat = describe(values)
    return f"{stat['median']:.1f} ({stat['iqr']:.1f})"


def sex_summary(group: pd.DataFrame) -> dict[str, str]:
    if "sex" not in group.columns:
        return {}
    sex = group.sex.fillna("Unknown").astype(str).replace("", "Unknown").str.title()
    total = len(sex)
    result = {}
    for label in ("Male", "Female", "Unknown"):
        count = int(sex.eq(label).sum())
        result[f"{label}, n (%)"] = f"{count} ({100.0 * count / total:.1f}%)"
    return result


def paper_table_1(case: pd.DataFrame, normalized: bool = False) -> pd.DataFrame:
    metrics = CENTRAL_METRICS
    if normalized:
        metrics = [f"{metric}_per_L_cuberoot" for metric in metrics]
    labels = [
        "Full Segmented Tracheal Length (mm)",
        "Tracheal EDT Mean Diameter (mm)",
        "LMB EDT Mean Diameter (mm)",
        "RMB EDT Mean Diameter (mm)",
    ]
    if normalized:
        labels = [label.replace("(mm)", "(mm/L^(1/3))") for label in labels]
    rows = []
    for bucket in AGE_ORDER + ["All"]:
        group = case if bucket == "All" else case[case.age_bucket == bucket]
        row = {
            "Age Range (y)": AGE_DISPLAY[bucket],
            "No. of Subjects": int(group.subject.nunique()),
            "No. of CT Scans": int(group.case_id.nunique()),
            "Age (y)": fmt_age(group.age_years),
        }
        row.update(sex_summary(group))
        for metric, label in zip(metrics, labels):
            row[label] = fmt_median_quartiles(group[metric])
        rows.append(row)
    return pd.DataFrame(rows)


def paper_table_2(case: pd.DataFrame, normalized: bool = False) -> pd.DataFrame:
    metrics = SECTION_METRICS
    if normalized:
        metrics = [f"{metric}_per_L_cuberoot" for metric in metrics]
    suffix = " (mm/L^(1/3))" if normalized else " (mm)"
    labels = ["Distal" + suffix, "Middle" + suffix, "Proximal" + suffix]
    rows = []
    for bucket in AGE_ORDER + ["All"]:
        group = case if bucket == "All" else case[case.age_bucket == bucket]
        row = {
            "Age Range (y)": AGE_DISPLAY[bucket],
            "No. of Subjects": int(group.subject.nunique()),
            "No. of CT Scans": int(group.case_id.nunique()),
        }
        row.update(sex_summary(group))
        for metric, label in zip(metrics, labels):
            row[label] = fmt_median_iqr(group[metric])
        rows.append(row)
    return pd.DataFrame(rows)


def add_branch_metadata(branches: pd.DataFrame, case: pd.DataFrame) -> pd.DataFrame:
    metadata = case[
        ["case_id", "subject", "age_years", "age_bucket", "total_segmented_lung_volume_mL", "lung_linear_scale_L_cuberoot"]
    ].drop_duplicates("case_id")
    data = branches.drop(columns=["subject"], errors="ignore").merge(
        metadata, on="case_id", validate="many_to_one"
    )
    data["diameter_mean_mm"] = 2.0 * data.radius_mean_mm
    data["radius_mean_mm_per_L_cuberoot"] = (
        data.radius_mean_mm / data.lung_linear_scale_L_cuberoot
    )
    data["diameter_mean_mm_per_L_cuberoot"] = (
        data.diameter_mean_mm / data.lung_linear_scale_L_cuberoot
    )
    data["centerline_length_mm_per_L_cuberoot"] = (
        data.centerline_length_mm / data.lung_linear_scale_L_cuberoot
    )
    data["branch_volume_fraction_ppm"] = (
        data.volume_mm3 / (data.total_segmented_lung_volume_mL * 1000.0) * 1e6
    )
    return data


def generation_tables(branches: pd.DataFrame):
    if MAX_REPORTED_GENERATION is not None:
        branches = branches[branches.generation <= MAX_REPORTED_GENERATION].copy()
    per_case = branches.groupby(
        ["age_bucket", "case_id", "subject", "generation"], as_index=False
    ).agg(
        branch_count=("branch_instance_id", "size"),
        mean_radius_mm=("radius_mean_mm", "mean"),
        median_radius_mm=("radius_mean_mm", "median"),
        mean_diameter_mm=("diameter_mean_mm", "mean"),
        median_diameter_mm=("diameter_mean_mm", "median"),
        mean_centerline_length_mm=("centerline_length_mm", "mean"),
        median_centerline_length_mm=("centerline_length_mm", "median"),
        mean_tortuosity=("tortuosity_length_ratio", "mean"),
        mean_branch_angle_deg=("branch_angle_deg", "mean"),
        total_parsed_branch_volume_mm3=("volume_mm3", "sum"),
        mean_radius_mm_per_L_cuberoot=("radius_mean_mm_per_L_cuberoot", "mean"),
        mean_diameter_mm_per_L_cuberoot=("diameter_mean_mm_per_L_cuberoot", "mean"),
        mean_centerline_length_mm_per_L_cuberoot=("centerline_length_mm_per_L_cuberoot", "mean"),
        total_branch_volume_fraction_ppm=("branch_volume_fraction_ppm", "sum"),
    )
    metrics = [
        "branch_count",
        "mean_radius_mm",
        "median_radius_mm",
        "mean_diameter_mm",
        "median_diameter_mm",
        "mean_centerline_length_mm",
        "median_centerline_length_mm",
        "mean_tortuosity",
        "mean_branch_angle_deg",
        "total_parsed_branch_volume_mm3",
        "mean_radius_mm_per_L_cuberoot",
        "mean_diameter_mm_per_L_cuberoot",
        "mean_centerline_length_mm_per_L_cuberoot",
        "total_branch_volume_fraction_ppm",
    ]
    rows = []
    maximum_generation = int(branches.generation.max())
    case_roster = branches[["age_bucket", "case_id", "subject"]].drop_duplicates()
    for bucket in AGE_ORDER:
        bucket_data = per_case[per_case.age_bucket == bucket]
        bucket_cases = case_roster[case_roster.age_bucket == bucket]
        for generation in range(1, maximum_generation + 1):
            group = bucket_data[bucket_data.generation == generation]
            all_scan_counts = bucket_cases.merge(
                group[["case_id", "branch_count"]], on="case_id", how="left"
            ).branch_count.fillna(0)
            count_all_stats = describe(all_scan_counts)
            row = {
                "age_group": AGE_DISPLAY[bucket],
                "age_bucket_code": bucket,
                "AirMorph_generation": int(generation),
                "n_cases_with_generation": int(group.case_id.nunique()),
                "n_unique_subjects_with_generation": int(group.subject.nunique()),
                "total_branches": int(group.branch_count.sum()),
                "n_scans_in_age_group": int(bucket_cases.case_id.nunique()),
                "branch_count_all_scans_mean": count_all_stats["mean"],
                "branch_count_all_scans_sd": count_all_stats["sd"],
                "branch_count_all_scans_median": count_all_stats["median"],
                "branch_count_all_scans_q1": count_all_stats["q1"],
                "branch_count_all_scans_q3": count_all_stats["q3"],
                "branch_count_all_scans_iqr": count_all_stats["iqr"],
            }
            for metric in metrics:
                stat = describe(group[metric])
                for name in ("mean", "sd", "median", "q1", "q3", "iqr", "min", "max"):
                    row[f"{metric}_{name}"] = stat[name]
            rows.append(row)
    return per_case, pd.DataFrame(rows)


def colors():
    cmap = plt.get_cmap("viridis")
    return {bucket: cmap(i / (len(AGE_ORDER) - 1)) for i, bucket in enumerate(AGE_ORDER)}


def plot_generation_dimensions(summary: pd.DataFrame, out: Path, normalized=False):
    palette = colors()
    if normalized:
        specs = [
            ("mean_radius_mm_per_L_cuberoot", "Radius / lung-volume linear scale (mm/L$^{1/3}$)"),
            ("mean_diameter_mm_per_L_cuberoot", "Diameter / lung-volume linear scale (mm/L$^{1/3}$)"),
        ]
        title = "Lung-volume-normalized airway dimensions by age group and AirMorph generation"
    else:
        specs = [
            ("mean_radius_mm", "EDT-derived lumen radius (mm)"),
            ("mean_diameter_mm", "EDT-derived lumen diameter (mm)"),
        ]
        title = "Airway dimensions by age group and AirMorph generation"
    fig, axes = plt.subplots(2, 1, figsize=(15, 11), sharex=True)
    for ax, (metric, ylabel) in zip(axes, specs):
        for bucket in AGE_ORDER:
            group = summary[summary.age_bucket_code == bucket].sort_values("AirMorph_generation")
            x = group.AirMorph_generation.to_numpy()
            y = group[f"{metric}_median"].to_numpy()
            sd = group[f"{metric}_sd"].fillna(0).to_numpy()
            ax.errorbar(
                x, y, yerr=sd, color=palette[bucket], marker="o", markersize=3.8,
                linewidth=1.5, elinewidth=0.8, capsize=2,
                label=f"{AGE_DISPLAY[bucket]} years",
            )
        ax.set_ylabel(ylabel + "\npoint: median; bars: ± between-scan SD")
        ax.set_ylim(bottom=0)
        ax.grid(True, alpha=0.25)
    max_generation = int(summary.AirMorph_generation.max())
    axes[-1].set_xlabel("AirMorph topologic generation (tracheal root = 1)")
    axes[-1].set_xticks(np.arange(1, max_generation + 1))
    axes[-1].set_xlim(0.5, max_generation + 0.5)
    axes[0].legend(ncol=4, fontsize=9, loc="upper right")
    fig.suptitle(title, fontsize=16)
    fig.text(
        0.5, 0.01,
        "All detected generations are shown. Estimates at high generations are based on fewer scans and are increasingly sensitive to CT resolution.",
        ha="center", fontsize=9,
    )
    fig.tight_layout(rect=(0, 0.035, 1, 0.97))
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_generation_counts(summary: pd.DataFrame, out: Path):
    palette = colors()
    fig, ax = plt.subplots(figsize=(15, 7))
    for bucket in AGE_ORDER:
        group = summary[summary.age_bucket_code == bucket].sort_values("AirMorph_generation")
        ax.errorbar(
            group.AirMorph_generation,
            group.branch_count_all_scans_mean,
            yerr=group.branch_count_all_scans_sd.fillna(0),
            color=palette[bucket], marker="o", markersize=4, linewidth=1.6,
            elinewidth=0.9, capsize=2, label=f"{AGE_DISPLAY[bucket]} years",
        )
    max_generation = int(summary.AirMorph_generation.max())
    ax.set_xlabel("AirMorph topologic generation (tracheal root = 1)")
    ax.set_ylabel("Branches per scan (mean ± SD)")
    ax.set_xticks(np.arange(1, max_generation + 1))
    ax.set_xlim(0.5, max_generation + 0.5)
    ax.set_ylim(bottom=0)
    ax.grid(True, alpha=0.25)
    ax.legend(ncol=4, fontsize=9, loc="upper right")
    ax.set_title("Number of detected airway branches by age group and generation")
    fig.tight_layout()
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)


def age_axis_stats(case: pd.DataFrame, metrics: list[str]) -> pd.DataFrame:
    rows = []
    for bucket in AGE_ORDER:
        group = case[case.age_bucket == bucket]
        row = {
            "bucket": bucket,
            "label": AGE_DISPLAY[bucket],
            "x": float(group.age_years.median()),
        }
        for metric in metrics:
            stat = describe(group[metric])
            row[f"{metric}_mean"] = stat["mean"]
            row[f"{metric}_sd"] = stat["sd"]
            row[f"{metric}_median"] = stat["median"]
        rows.append(row)
    return pd.DataFrame(rows)


def plot_central(case: pd.DataFrame, out: Path, normalized=False):
    metrics = CENTRAL_METRICS
    if normalized:
        metrics = [f"{metric}_per_L_cuberoot" for metric in metrics]
    stats = age_axis_stats(case, metrics)
    fig, axes = plt.subplots(2, 1, figsize=(12, 10), sharex=True)
    length = metrics[0]
    axes[0].errorbar(stats.x, stats[f"{length}_mean"], yerr=stats[f"{length}_sd"],
                     color="#264653", marker="o", linewidth=2, capsize=4)
    axes[0].set_ylabel("Full segmented tracheal length" +
                       (" (mm/L$^{1/3}$)" if normalized else " (mm)") + "\nmean ± SD")
    diameter_labels = ["Trachea", "LMB", "RMB"]
    diameter_colors = ["#2a9d8f", "#e9c46a", "#e76f51"]
    for metric, label, color in zip(metrics[1:], diameter_labels, diameter_colors):
        axes[1].errorbar(stats.x, stats[f"{metric}_mean"], yerr=stats[f"{metric}_sd"],
                         color=color, label=label, marker="o", linewidth=2, capsize=4)
    axes[1].set_ylabel("EDT-derived mean lumen diameter" +
                       (" (mm/L$^{1/3}$)" if normalized else " (mm)") + "\nmean ± SD")
    axes[1].set_xlabel("Median age within requested age group (years)")
    axes[1].legend()
    for ax in axes:
        ax.grid(True, alpha=0.25)
        ax.set_ylim(bottom=0)
    title = "Central-airway dimensions by age group"
    if normalized:
        title += " after lung-volume normalization"
    fig.suptitle(title, fontsize=15)
    fig.tight_layout()
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_sections(case: pd.DataFrame, out: Path, normalized=False):
    metrics = SECTION_METRICS
    if normalized:
        metrics = [f"{metric}_per_L_cuberoot" for metric in metrics]
    stats = age_axis_stats(case, metrics)
    fig, ax = plt.subplots(figsize=(12, 7))
    for metric, label, color in zip(metrics, ["Distal", "Middle", "Proximal"],
                                    ["#1b9e77", "#7570b3", "#d95f02"]):
        ax.errorbar(stats.x, stats[f"{metric}_mean"], yerr=stats[f"{metric}_sd"],
                    color=color, label=label, marker="o", linewidth=2, capsize=4)
    ax.set_xlabel("Median age within requested age group (years)")
    ax.set_ylabel("EDT-derived tracheal lumen diameter" +
                  (" (mm/L$^{1/3}$)" if normalized else " (mm)") + "\nmean ± SD")
    title = "Tracheal diameter by full-segmented-length section and age group"
    if normalized:
        title += " after lung-volume normalization"
    ax.set_title(title)
    ax.set_ylim(bottom=0)
    ax.grid(True, alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_lung_volume(case: pd.DataFrame, out: Path):
    stats = age_axis_stats(case, ["total_segmented_lung_volume_mL"])
    fig, ax = plt.subplots(figsize=(11, 6))
    metric = "total_segmented_lung_volume_mL"
    ax.errorbar(stats.x, stats[f"{metric}_mean"] / 1000.0,
                yerr=stats[f"{metric}_sd"] / 1000.0,
                marker="o", linewidth=2, capsize=4, color="#2878b5")
    ax.set_xlabel("Median age within requested age group (years)")
    ax.set_ylabel("Total segmented lung volume (L)\nmean ± SD")
    ax.set_title("Total segmented lung volume by age group")
    ax.set_ylim(bottom=0)
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    fig.savefig(out, dpi=300, bbox_inches="tight")
    plt.close(fig)


def style_workbook(path: Path):
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    dark, medium, light = "17365D", "366092", "DCE6F1"
    thin = Side(style="thin", color="A6A6A6")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=dark)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[1].height = 42
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        for column_cells in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column_cells[:150]]
            width = min(42, max(11, max(values, default=10) + 2))
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        if sheet.title.startswith("Table"):
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "All":
                    for cell in sheet[row]:
                        cell.fill = PatternFill("solid", fgColor=light)
                        cell.font = Font(bold=True)
    workbook.save(path)


def write_readme(out: Path, case: pd.DataFrame, summary: pd.DataFrame):
    counts = case.groupby("age_bucket").case_id.nunique().reindex(AGE_ORDER).fillna(0)
    count_text = ", ".join(f"{AGE_DISPLAY[b]}: n={int(counts[b])}" for b in AGE_ORDER)
    max_gen = int(summary.AirMorph_generation.max())
    populated = [bucket for bucket in AGE_ORDER if (case.age_bucket == bucket).any()]
    youngest_bucket, oldest_bucket = populated[0], populated[-1]
    youngest = case[case.age_bucket == youngest_bucket]
    oldest = case[case.age_bucket == oldest_bucket]
    deepest = summary[
        (summary.AirMorph_generation == max_gen)
        & (summary.n_cases_with_generation > 0)
    ]
    deepest_text = ", ".join(
        f"{row.age_group}: {int(row.n_cases_with_generation)}/{int(row.n_scans_in_age_group)} scans"
        for row in deepest.itertuples()
    )
    text = f"""# Normal pediatric airway results — radiologist/professor meeting

Prepared {date.today().isoformat()} from {len(case)} normal CT scans representing {case.subject.nunique()} unique anonymized subjects.

## Open these first

1. `airway_meeting_results.xlsx` — paper-style tables, complete descriptive statistics, per-generation results, lung-volume-normalized sensitivity analyses, case-level source values, methods, and QC.
2. `figures/01_generation_dimensions_median_with_sd.png` — closest available analogue to Fig. 4 in the second reference paper.
3. `figures/04_central_airway_dimensions_mean_sd.png` — tracheal length and central-airway diameters by age group.
4. `figures/06_tracheal_sections_mean_sd.png` — distal, middle, and proximal tracheal lumen-diameter surrogates.

## Descriptive meeting highlights

- Median full segmented tracheal length increases from {youngest.full_segmented_tracheal_length_mm.median():.1f} mm at {AGE_DISPLAY[youngest_bucket]} years to {oldest.full_segmented_tracheal_length_mm.median():.1f} mm at {AGE_DISPLAY[oldest_bucket]} years.
- Median EDT-derived tracheal diameter increases from {youngest.tracheal_edt_mean_diameter_mm.median():.1f} mm to {oldest.tracheal_edt_mean_diameter_mm.median():.1f} mm across those age groups.
- Median LMB/RMB diameter is {youngest.lmb_edt_mean_diameter_mm.median():.1f}/{youngest.rmb_edt_mean_diameter_mm.median():.1f} mm at {AGE_DISPLAY[youngest_bucket]} years and {oldest.lmb_edt_mean_diameter_mm.median():.1f}/{oldest.rmb_edt_mean_diameter_mm.median():.1f} mm at {AGE_DISPLAY[oldest_bucket]} years.
- The deepest retained AirMorph generation is {max_gen}; it is present in {deepest_text}. Treat this sparse tail as a detection result, not a biological age trend.

## Cohort and requested age groups

The requested bucket definitions are retained exactly from `14_normal_age_distribution.py`: [0,3), [3,6), [6,9], [10,12), [12,14), [14,16), and [16,18]. Counts: {count_text}.

There are repeated scans: {len(case)} scans arise from {case.subject.nunique()} subjects. Tables show both counts. Descriptive values use each CT scan as one observation and therefore do not establish independent-subject reference limits.

## Measurement terminology

- **Full segmented tracheal length** is the physical root-centerline path from the carina-side endpoint to the superior endpoint of the segmented airway. It includes any segmented extrathoracic trachea and is not the intrathoracic carina-to-thoracic-inlet length used by Kuo et al.
- **EDT-derived lumen radius** is the Euclidean distance to the nearest segmented airway boundary sampled on the branch centerline. Diameter is two times this radius. It is a lumen-caliber surrogate, not the mean of short and long perpendicular axes and not an outer-wall measurement.
- **Tracheal sections** divide the full segmented root path into equal distal, middle, and proximal thirds. Diameters are sampled at 1-mm centerline intervals and averaged within each third.
- **Generation** is AirMorph topologic generation with the tracheal root equal to 1. All detected generations 1–{max_gen} are retained. This numbering is different from the second paper, where segmental bronchi are G0.
- **Total segmented lung volume** is calculated from all nonzero voxels in AirMorph's `lunglobe.nii.gz`. Despite that filename, this is a whole-lung left/right mask generated by LungMask U-Net R231, not a five-lobe mask. Voxel count is multiplied by physical voxel volume and divided by 1000 to obtain mL.
- **Lung-volume normalization** is exploratory. Each length/radius/diameter is divided by the cube root of total segmented lung volume in liters within the same scan, then summarized by age group. This preserves dimensional scaling and is not a published pediatric reference standard.

## What cannot be reported from these outputs

Bronchial outer diameter, wall thickness, bronchial wall area, paired pulmonary-artery diameter, bronchus–artery ratios, iodine-contrast subgroup results, true perpendicular short/long axes, and a manually identified thoracic-inlet plane are unavailable. No values for these outcomes were imputed.

## Interpretation limits for the meeting

All anatomical labels are AirMorph predictions and require pediatric radiologist review. High-generation results are based on progressively fewer visible branches and scans and are increasingly sensitive to spatial resolution. The R231 whole-lung masks provide an automated total lung-volume surrogate; they do not prove inspiratory adequacy. Contrast status is unavailable, so the generation figures combine contrast and noncontrast CT scans.

## Reference papers

- Kuo W, et al. *Reference Values for Central Airway Dimensions on CT Images of Children and Adolescents.* AJR 2018. DOI: https://doi.org/10.2214/AJR.17.18597
- Lv Q, et al. *Normative values for lung, bronchial sizes, and bronchus-artery ratios in chest CT scans: from infancy into young adulthood.* European Radiology 2025. https://doi.org/10.1007/s00330-025-11367-w
"""
    (out / "README_FIRST.md").write_text(text)


def main() -> None:
    args = parse_args()
    out = args.out
    csv_dir = out / "csv"
    fig_dir = out / "figures"
    supporting = out / "supporting_files"
    for directory in (out, csv_dir, fig_dir, supporting):
        directory.mkdir(parents=True, exist_ok=True)

    branches = pd.read_csv(
        BRANCHES,
        low_memory=False,
        dtype={"branch_instance_id": "Int64", "parent_instance_id": "Int64"},
    )
    age = pd.read_csv(AGES)
    status = pd.read_csv(STATUS)
    cache = csv_dir / "imaging_measurement_cache.csv"
    case = build_case_table(branches, age, status, cache, args.reuse_case_cache)
    case = case.sort_values(["age_years", "case_id"]).reset_index(drop=True)

    branch_data = add_branch_metadata(branches, case)
    per_case_generation, generation_summary = generation_tables(branch_data)

    table1 = paper_table_1(case)
    table2 = paper_table_2(case)
    table1_norm = paper_table_1(case, normalized=True)
    table2_norm = paper_table_2(case, normalized=True)

    raw_metrics = [
        "age_years", "total_segmented_lung_volume_mL",
        *CENTRAL_METRICS, *SECTION_METRICS,
        "raw_root_chord_length_mm", "raw_root_edt_mean_diameter_mm",
    ]
    normalized_metrics = [
        f"{metric}_per_L_cuberoot" for metric in CENTRAL_METRICS + SECTION_METRICS
    ]
    central_numeric = grouped_descriptive(case, raw_metrics)
    normalized_numeric = grouped_descriptive(case, normalized_metrics)

    base_generation_columns = [
        "age_group", "age_bucket_code", "AirMorph_generation",
        "n_cases_with_generation", "n_unique_subjects_with_generation",
        "n_scans_in_age_group", "total_branches",
    ]
    generation_radius = generation_summary[base_generation_columns + [
        f"mean_radius_mm_{stat}"
        for stat in ("mean", "sd", "median", "q1", "q3", "iqr", "min", "max")
    ]]
    generation_diameter = generation_summary[base_generation_columns + [
        f"mean_diameter_mm_{stat}"
        for stat in ("mean", "sd", "median", "q1", "q3", "iqr", "min", "max")
    ]]
    generation_counts = generation_summary[base_generation_columns + [
        f"branch_count_all_scans_{stat}"
        for stat in ("mean", "sd", "median", "q1", "q3", "iqr")
    ]]
    generation_normalized = generation_summary[base_generation_columns + [
        f"mean_diameter_mm_per_L_cuberoot_{stat}"
        for stat in ("mean", "sd", "median", "q1", "q3", "iqr", "min", "max")
    ]]

    case.to_csv(csv_dir / "case_level_central_airway_and_lung_volume.csv", index=False)
    per_case_generation.to_csv(csv_dir / "per_scan_generation_metrics.csv", index=False)
    generation_summary.to_csv(csv_dir / "age_group_generation_complete_statistics.csv", index=False)
    central_numeric.to_csv(csv_dir / "central_airway_complete_statistics.csv", index=False)
    normalized_numeric.to_csv(csv_dir / "lung_normalized_complete_statistics.csv", index=False)
    table1.to_csv(csv_dir / "table_1_central_airway_by_age.csv", index=False)
    table2.to_csv(csv_dir / "table_2_tracheal_sections_by_age.csv", index=False)
    table1_norm.to_csv(csv_dir / "table_1_lung_normalized.csv", index=False)
    table2_norm.to_csv(csv_dir / "table_2_lung_normalized.csv", index=False)

    methods = pd.DataFrame(
        [
            ("Age groups", "; ".join(AGE_DISPLAY[bucket] for bucket in AGE_ORDER)),
            ("Central airway identification", "Single graph root = trachea. The two root children with the largest descendant trees = main bronchi; side assigned from majority downstream left/right anatomical-label votes."),
            ("Full segmented tracheal length", "Longest physical root tree path, carina-side endpoint to superior segmented endpoint; includes segmented extrathoracic trachea."),
            ("Radius and diameter", "Physical Euclidean distance-transform radius sampled on branch centerline; diameter = 2 × radius. Lumen surrogate only."),
            ("Tracheal sections", "Full root path oriented carina to superior endpoint, sampled every 1 mm, divided into equal distal/middle/proximal thirds; within-scan section value is mean sampled diameter."),
            ("Generation summary", "Branches aggregated within scan and generation first; scans then summarized within age group. " + (f"Reported generations are capped at {MAX_REPORTED_GENERATION}." if MAX_REPORTED_GENERATION is not None else "All AirMorph generations retained.")),
            ("Graph error bars", "Point is age-group median of per-scan generation means; bars show ± between-scan SD, as requested."),
            ("Normalization", "Within each scan, metric in mm divided by cube root of total segmented lung volume in liters; group summaries computed afterward."),
            ("Lung mask provenance", "AirMorph lunglobe.nii.gz generated by its embedded LungMask U-Net R231 model (left/right whole-lung labels, not five-lobe segmentation). Total volume = nonzero voxels × physical voxel volume / 1000 mL."),
            ("Repeated scans", f"{len(case)} scans from {case.subject.nunique()} cohort identifiers. Descriptive tables report both counts."),
            ("Unavailable outcomes", "No outer diameter, wall thickness, arterial pairing, contrast status, true cross-sectional axes, or manual thoracic-inlet landmark."),
        ],
        columns=["Item", "Definition"],
    )
    qc = pd.DataFrame(
        [
            ("Normal CT scans", len(case)),
            ("Unique subjects", case.subject.nunique()),
            ("Branch rows", len(branch_data)),
            ("Minimum AirMorph generation", int(branch_data.generation.min())),
            ("Maximum reported AirMorph generation", int(min(branch_data.generation.max(), MAX_REPORTED_GENERATION)) if MAX_REPORTED_GENERATION is not None else int(branch_data.generation.max())),
            ("Cases with one central root and opposed L/R main bronchi", len(case)),
            ("Minimum LMB side-vote margin", int(case.lmb_side_vote_margin.min())),
            ("Minimum RMB side-vote margin", int(case.rmb_side_vote_margin.min())),
            ("Missing requested age buckets", int(case.age_bucket.isna().sum())),
            ("Missing total lung volumes", int(case.total_segmented_lung_volume_mL.isna().sum())),
            ("Maximum absolute root path-vs-exported MST difference (mm)", float(case.trachea_path_vs_raw_mst_difference_mm.abs().max())),
            ("Root paths with >2 MST endpoints", int((case.root_tree_endpoint_count > 2).sum())),
        ],
        columns=["QC check", "Result"],
    )

    workbook = out / "airway_meeting_results.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        table1.to_excel(writer, sheet_name="Table 1 Central", index=False)
        table2.to_excel(writer, sheet_name="Table 2 Tracheal Sections", index=False)
        table1_norm.to_excel(writer, sheet_name="Table 1 Normalized", index=False)
        table2_norm.to_excel(writer, sheet_name="Table 2 Sections Normalized", index=False)
        central_numeric.to_excel(writer, sheet_name="Central Complete Stats", index=False)
        normalized_numeric.to_excel(writer, sheet_name="Normalized Complete Stats", index=False)
        generation_radius.to_excel(writer, sheet_name="Generation Radius", index=False)
        generation_diameter.to_excel(writer, sheet_name="Generation Diameter", index=False)
        generation_counts.to_excel(writer, sheet_name="Generation Counts", index=False)
        generation_normalized.to_excel(writer, sheet_name="Generation Normalized", index=False)
        generation_summary.to_excel(writer, sheet_name="Generation Complete Stats", index=False)
        per_case_generation.to_excel(writer, sheet_name="Per Scan Generation", index=False)
        case.drop(columns="case_dir").to_excel(writer, sheet_name="Case Level", index=False)
        methods.to_excel(writer, sheet_name="Methods", index=False)
        qc.to_excel(writer, sheet_name="QC", index=False)
    style_workbook(workbook)

    plot_generation_dimensions(generation_summary, fig_dir / "01_generation_dimensions_median_with_sd.png")
    plot_generation_counts(generation_summary, fig_dir / "02_generation_branch_counts_mean_sd.png")
    plot_generation_dimensions(generation_summary, fig_dir / "03_generation_dimensions_lung_normalized.png", normalized=True)
    plot_central(case, fig_dir / "04_central_airway_dimensions_mean_sd.png")
    plot_central(case, fig_dir / "05_central_airway_dimensions_lung_normalized.png", normalized=True)
    plot_sections(case, fig_dir / "06_tracheal_sections_mean_sd.png")
    plot_sections(case, fig_dir / "07_tracheal_sections_lung_normalized.png", normalized=True)
    plot_lung_volume(case, fig_dir / "08_total_segmented_lung_volume_mean_sd.png")

    for source in [
        RUN / "radiologist_qc" / "radiologist_review.xlsx",
        RUN / "radiologist_qc" / "tables" / "analysis_ready_airway_branch_morphometry.xlsx",
        RUN / "normal_cases_age_distribution.txt",
    ]:
        if source.exists():
            shutil.copy2(source, supporting / source.name)
    representative = [
        "sub-0012_acc-0295", "sub-0010_acc-0121", "sub-0253_acc-0421"
    ]
    rep_dir = supporting / "representative_case_profiles"
    rep_dir.mkdir(exist_ok=True)
    for case_id in representative:
        source = RUN / "generation_profiles" / "cases" / f"{case_id}_four_panel_generation_profile.png"
        if source.exists():
            shutil.copy2(source, rep_dir / source.name)

    write_readme(out, case, generation_summary)

    archive = out.with_suffix(".tar.gz")
    with tarfile.open(archive, "w:gz") as tar:
        tar.add(out, arcname=out.name)
    print(f"wrote meeting folder: {out}")
    print(f"wrote archive: {archive}")


if __name__ == "__main__":
    main()
